import datetime
import io
from pathlib import Path

from fastapi.testclient import TestClient
import openpyxl
import pytest

from backend.api import app

client = TestClient(app)

RAIZ = Path(__file__).resolve().parent.parent.parent

# Header de autenticação padrão para testes
AUTH_HEADERS = {"Authorization": "Bearer dev-token-contador@bregalda.com.br"}


def _criar_xlsx_teste_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Razão"

    # Cabeçalho contábil
    ws.append(["Empresa:", "COMERCIO DE BEBIDAS LTDA", "", ""])
    ws.append(["CNPJ:", "12.345.678/0001-99", "", ""])
    ws.append(["Período:", "01/07/2026 a 31/07/2026", "", ""])
    ws.append([])

    # Linha de títulos
    ws.append(["Data", "Histórico", "Cta.C.Part.", "Débito", "Crédito", "Saldo"])

    # Conta
    ws.append(["Conta: 504  2.1.30.100.1  DUPLICATAS A PAGAR", "", "", "", "", ""])

    # Saldo anterior
    ws.append(["", "SALDO ANTERIOR", "", "", "", 10500.00])

    # Lançamentos
    ws.append(["01/07/2026", "NF 1001 CERVEJARIA BRASIL", "1.1.20", "", 1500.00, 12000.00])
    ws.append([datetime.date(2026, 7, 10), "PAGTO CERVEJARIA BRASIL", "1.1.10", 1500.00, "", 10500.00])
    ws.append(["15/07/2026", "NF 1002 DISTRIBUIDORA VALE", "1.1.20", "", 850.50, 11350.50])
    ws.append(["20/07/2026", "PAGTO ANTIGO DIVERSO", "1.1.10", 300.00, "", 11050.50])

    # Total do mês
    ws.append(["", "Total do mês", "", 1800.00, 2350.50, ""])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_health():
    response = client.get("/health")
    assert response.status_code == 200
    dados = response.json()
    assert dados == {"status": "ok", "app": "conferidor-razao"}


def test_processar_xlsx():
    xlsx_bytes = _criar_xlsx_teste_bytes()
    response = client.post(
        "/api/processar",
        headers=AUTH_HEADERS,
        files={"file": ("planilha_teste.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    dados = response.json()

    assert dados["arquivo"] == "planilha_teste.xlsx"
    assert "COMERCIO DE BEBIDAS LTDA" in dados["empresa"]
    assert "12.345.678/0001-99" in dados["cnpj"]
    assert dados["conta"] == "2.1.30.100.1"
    assert len(dados["lancamentos"]) == 4

    resumo = dados["resumo"]
    assert resumo["creditos"]["qtd"] == 2
    assert resumo["debitos"]["qtd"] == 2
    assert resumo["quitados"]["qtd"] == 1
    assert resumo["abertos"]["qtd"] == 1
    assert resumo["sem_par"]["qtd"] == 1
    assert resumo["cobertura_qtd"] == 50.0


def test_processar_xls_real():
    caminhos = [
        RAIZ / "docs" / "internos" / "Razão (1).xls",
        RAIZ / "Razão (1).xls",
    ]
    caminho_xls = next((c for c in caminhos if c.exists()), None)
    if not caminho_xls:
        pytest.skip("Planilha real mantida em docs/internos/ (sigilo/gitignored)")

    with open(caminho_xls, "rb") as f:
        xls_bytes = f.read()

    response = client.post(
        "/api/processar",
        headers=AUTH_HEADERS,
        files={"file": ("Razão (1).xls", xls_bytes, "application/vnd.ms-excel")},
    )
    assert response.status_code == 200
    dados = response.json()

    assert dados["arquivo"] == "Razão (1).xls"
    assert len(dados["lancamentos"]) > 100
    assert "resumo" in dados
    assert dados["resumo"]["cobertura_qtd"] > 80.0
    assert dados["resumo"]["quitados"]["qtd"] > 400


def test_processar_formato_invalido():
    response = client.post(
        "/api/processar",
        headers=AUTH_HEADERS,
        files={"file": ("relatorio.pdf", b"%PDF-1.4 ...", "application/pdf")},
    )
    assert response.status_code == 400
    assert "Formato não suportado" in response.json()["detail"]


def test_processar_xlsx_corrompido():
    response = client.post(
        "/api/processar",
        headers=AUTH_HEADERS,
        files={"file": ("corrompido.xlsx", b"isto nao e um zip xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 422


def test_processar_xls_corrompido():
    response = client.post(
        "/api/processar",
        headers=AUTH_HEADERS,
        files={"file": ("corrompido.xls", b"isto nao e um ole xls", "application/vnd.ms-excel")},
    )
    assert response.status_code == 422


def test_exportar_excel():
    # Primeiro processa uma planilha válida para obter a estrutura completa do JSON
    xlsx_bytes = _criar_xlsx_teste_bytes()
    proc_response = client.post(
        "/api/processar",
        headers=AUTH_HEADERS,
        files={"file": ("conferencia_teste.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert proc_response.status_code == 200
    razao_dict = proc_response.json()

    # Chama o endpoint de exportação
    exp_response = client.post(
        "/api/exportar-excel",
        headers=AUTH_HEADERS,
        json=razao_dict,
    )
    assert exp_response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in exp_response.headers["content-type"]
    assert 'attachment; filename="conferencia_conferencia_teste.xlsx"' in exp_response.headers["content-disposition"]

    # Valida que o conteúdo retornado é um arquivo .xlsx válido e legível por openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(exp_response.content))
    assert "Razão Conciliado" in wb.sheetnames
    ws = wb["Razão Conciliado"]
    assert ws.max_row > 5


def test_auth_login_dev_sucesso():
    response = client.post(
        "/api/auth/login",
        json={"email": "contador@bregalda.com.br", "password": "qualquer_senha"},
    )
    assert response.status_code == 200
    dados = response.json()
    assert "usuario" in dados
    assert dados["usuario"]["email"] == "contador@bregalda.com.br"
    assert "token" in dados
    assert dados["token"].startswith("dev-token-")


def test_auth_login_faltando_campos():
    response = client.post(
        "/api/auth/login",
        json={"email": "", "password": ""},
    )
    assert response.status_code == 400
    assert "informe seu e-mail" in response.json()["detail"]


def test_auth_me_com_sucesso():
    response = client.get("/api/auth/me", headers=AUTH_HEADERS)
    assert response.status_code == 200
    dados = response.json()
    assert "usuario" in dados
    assert dados["usuario"]["email"] == "contador@bregalda.com.br"


def test_auth_me_sem_token():
    response = client.get("/api/auth/me")
    assert response.status_code == 401


def test_auth_logout_sucesso():
    response = client.post("/api/auth/logout", headers=AUTH_HEADERS)
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_auth_login_ricardo_sucesso():
    response = client.post(
        "/api/auth/login",
        json={"email": "ricardo@bregalda.com.br", "password": "bregalda2026"},
    )
    assert response.status_code == 200
    dados = response.json()
    assert dados["usuario"]["email"] == "ricardo@bregalda.com.br"
    assert dados["usuario"]["nome"] == "Ricardo Bregalda"
    assert dados["token"] == "dev-token-ricardo@bregalda.com.br"


def test_auth_login_ricardo_senha_incorreta():
    response = client.post(
        "/api/auth/login",
        json={"email": "ricardo@bregalda.com.br", "password": "senha_errada"},
    )
    assert response.status_code == 401
    assert "Senha incorreta" in response.json()["detail"]


