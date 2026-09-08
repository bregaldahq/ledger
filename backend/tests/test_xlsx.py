import io
import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill
import pytest

from backend.ledger import Razao, Lancamento
from backend.reconciler import conciliar
from backend.xls_reader import ArquivoInvalido
from backend.xlsx_reader import ler as ler_xlsx, ler_xlsx as ler_xlsx_alias
from backend.xlsx_exporter import gerar_xlsx_colorido


def _criar_xlsx_exemplo_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Razão"

    # Cabeçalho contábil
    ws.append(["Empresa:", "COMERCIO DE BEBIDAS LTDA", "", ""])
    ws.append(["CNPJ:", "12.345.678/0001-99", "", ""])
    ws.append(["Período:", "01/07/2026 a 31/07/2026", "", ""])
    ws.append([])  # linha em branco

    # Linha de títulos
    ws.append(["Data", "Histórico", "Cta.C.Part.", "Débito", "Crédito", "Saldo"])

    # Conta
    ws.append(["Conta: 504  2.1.30.100.1  DUPLICATAS A PAGAR", "", "", "", "", ""])

    # Saldo anterior
    ws.append(["", "SALDO ANTERIOR", "", "", "", 10500.00])

    # Lançamentos
    # Crédito (aquisição)
    ws.append(["01/07/2026", "NF 1001 CERVEJARIA BRASIL", "1.1.20", "", 1500.00, 12000.00])
    # Débito (pagamento)
    ws.append([datetime.date(2026, 7, 10), "PAGTO CERVEJARIA BRASIL", "1.1.10", 1500.00, "", 10500.00])
    # Crédito em aberto
    ws.append(["15/07/2026", "NF 1002 DISTRIBUIDORA VALE", "1.1.20", "", 850.50, 11350.50])
    # Débito sem par
    ws.append(["20/07/2026", "PAGTO ANTIGO DIVERSO", "1.1.10", 300.00, "", 11050.50])

    # Total do mês
    ws.append(["", "Total do mês", "", 1800.00, 2350.50, ""])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_leitura_xlsx_em_memoria():
    conteudo = _criar_xlsx_exemplo_bytes()
    razao = ler_xlsx(conteudo, "razao_teste.xlsx")

    assert isinstance(razao, Razao)
    assert razao.arquivo == "razao_teste.xlsx"
    assert "COMERCIO DE BEBIDAS LTDA" in razao.empresa
    assert "12.345.678/0001-99" in razao.cnpj
    assert "01/07/2026 a 31/07/2026" in razao.periodo
    assert razao.conta == "2.1.30.100.1"
    assert "DUPLICATAS A PAGAR" in razao.conta_nome
    assert razao.saldo_anterior == 1050000
    assert razao.total_debito == 180000
    assert razao.total_credito == 235050
    assert razao.colunas_ordenadas is False

    # 4 lançamentos: 2 créditos, 2 débitos
    assert len(razao.lancamentos) == 4
    c1, d1, c2, d2 = razao.lancamentos

    assert c1.tipo == "C"
    assert c1.valor == 150000
    assert c1.data == "2026-07-01"
    assert "CERVEJARIA BRASIL" in c1.historico

    assert d1.tipo == "D"
    assert d1.valor == 150000
    assert d1.data == "2026-07-10"

    assert c2.tipo == "C"
    assert c2.valor == 85050
    assert c2.data == "2026-07-15"

    assert d2.tipo == "D"
    assert d2.valor == 30000
    assert d2.data == "2026-07-20"

    # Conciliação
    resumo = conciliar(razao)
    assert resumo["quitados"]["qtd"] == 1
    assert resumo["quitados"]["valor"] == 150000
    assert resumo["abertos"]["qtd"] == 1
    assert resumo["sem_par"]["qtd"] == 1
    assert c1.status == "conciliado"
    assert d1.status == "conciliado"
    assert c1.par == d1.id
    assert d1.par == c1.id
    assert c1.dias == 9


def test_alias_ler_xlsx():
    conteudo = _criar_xlsx_exemplo_bytes()
    razao = ler_xlsx_alias(conteudo, "alias.xlsx")
    assert razao.arquivo == "alias.xlsx"
    assert len(razao.lancamentos) == 4


def test_leitura_xlsx_caminho_disco(tmp_path: Path):
    conteudo = _criar_xlsx_exemplo_bytes()
    caminho = tmp_path / "planilha.xlsx"
    caminho.write_bytes(conteudo)

    razao = ler_xlsx(str(caminho), "planilha.xlsx")
    assert razao.arquivo == "planilha.xlsx"
    assert len(razao.lancamentos) == 4
    assert razao.conta == "2.1.30.100.1"


def test_leitura_xlsx_variacoes_data():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Data", "Histórico", "Cta.C.Part.", "Débito", "Crédito", "Saldo"])
    ws.append(["Conta: 2.1.30.100.1 FORNECEDORES", "", "", "", "", ""])
    # datetime.date
    ws.append([datetime.date(2026, 7, 5), "DOC DATE", "", 100.0, "", ""])
    # datetime.datetime
    ws.append([datetime.datetime(2026, 7, 6, 12, 0), "DOC DATETIME", "", "", 100.0, ""])
    # String YYYY-MM-DD
    ws.append(["2026-07-07", "DOC ISO", "", 200.0, "", ""])
    # String DD/MM/YYYY
    ws.append(["08/07/2026", "DOC BR", "", "", 200.0, ""])

    buf = io.BytesIO()
    wb.save(buf)
    razao = ler_xlsx(buf.getvalue(), "datas.xlsx")

    assert [l.data for l in razao.lancamentos] == [
        "2026-07-05",
        "2026-07-06",
        "2026-07-07",
        "2026-07-08",
    ]


def test_leitura_xlsx_invalido_sem_colunas():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["A", "B", "C"])
    ws.append([1, 2, 3])
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(ArquivoInvalido) as exc:
        ler_xlsx(buf.getvalue(), "invalido.xlsx")
    assert "Débito e Crédito" in str(exc.value)


def test_leitura_xlsx_vazio():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.delete_rows(1, 10)
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(ArquivoInvalido):
        ler_xlsx(buf.getvalue(), "vazio.xlsx")


def test_exportacao_xlsx_colorido_basico():
    razao_dict = {
        "empresa": "INDUSTRIA EXEMPLO S/A",
        "cnpj": "98.765.432/0001-10",
        "periodo": "01/08/2026 a 31/08/2026",
        "conta": "2.1.20.001",
        "conta_nome": "FORNECEDORES DIVERSOS",
        "saldo_anterior": 500000,
        "total_debito": 150000,
        "total_credito": 250000,
        "colunas_ordenadas": False,
        "lancamentos": [
            {
                "id": 0,
                "linha": 1,
                "tipo": "C",
                "valor": 100000,
                "data": "2026-08-01",
                "historico": "COMPRA A",
                "contrapartida": "1.1.01",
                "status": "quitado",
                "par": 1,
                "dias": 5,
            },
            {
                "id": 1,
                "linha": 2,
                "tipo": "D",
                "valor": 100000,
                "data": "2026-08-06",
                "historico": "PAGAMENTO A",
                "contrapartida": "1.1.02",
                "status": "quitado",
                "par": 0,
                "dias": 5,
            },
            {
                "id": 2,
                "linha": 3,
                "tipo": "C",
                "valor": 150000,
                "data": "2026-08-10",
                "historico": "COMPRA B",
                "contrapartida": "1.1.01",
                "status": "aberto",
                "par": None,
                "dias": None,
            },
            {
                "id": 3,
                "linha": 4,
                "tipo": "D",
                "valor": 50000,
                "data": "2026-08-20",
                "historico": "PAGTO ANTERIOR",
                "contrapartida": "1.1.02",
                "status": "sem_par",
                "par": None,
                "dias": None,
            },
        ],
        "resumo": {
            "cobertura_valor": 40.0,
            "quitados": {"qtd": 1, "valor": 100000},
            "abertos": {"qtd": 1, "valor": 150000},
            "sem_par": {"qtd": 1, "valor": 50000},
        },
    }

    resultado_bytes = gerar_xlsx_colorido(razao_dict)
    assert isinstance(resultado_bytes, bytes)
    assert len(resultado_bytes) > 0

    # Abrir e inspecionar a planilha gerada
    wb = openpyxl.load_workbook(io.BytesIO(resultado_bytes))
    ws = wb.active

    # Verificar cabeçalho contábil presente
    textos = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("INDUSTRIA EXEMPLO S/A" in t for t in textos)
    assert any("98.765.432/0001-10" in t for t in textos)
    assert any("01/08/2026 a 31/08/2026" in t for t in textos)

    # Localizar cabeçalho da tabela
    linha_titulos = None
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and "Débito" in row and "Crédito" in row:
            linha_titulos = r_idx
            break
    assert linha_titulos is not None

    # Verificar linhas de dados após os títulos
    status_encontrados = []
    cores_encontradas = []

    for r_idx in range(linha_titulos + 1, ws.max_row + 1):
        row_cells = [ws.cell(r_idx, c_idx) for c_idx in range(1, ws.max_column + 1)]
        status_cell = row_cells[6]  # 7ª coluna: Status / Conciliação
        if status_cell.value:
            status_encontrados.append(str(status_cell.value))
            # Verificar se a célula ou linha possui preenchimento colorido
            fill = status_cell.fill
            if fill and fill.start_color and fill.start_color.rgb:
                cores_encontradas.append(fill.start_color.rgb.upper())

    # Status esperados
    assert any("Quitado em 5 dias (par #1)" in s for s in status_encontrados)
    assert any("Em aberto" in s for s in status_encontrados)
    assert any("Sem par no mês" in s for s in status_encontrados)

    # Cores esperadas (amarelo claro FEF9C3, vermelho claro FEE2E2, azul claro DBEAFE)
    # openpyxl pode prefixar com FF (ARGB: FFFEF9C3, FFFEE2E2, FFDBEAFE)
    cores_hex = [c[-6:] for c in cores_encontradas]
    assert "FEF9C3" in cores_hex  # Quitado
    assert "FEE2E2" in cores_hex  # Aberto
    assert "DBEAFE" in cores_hex  # Sem par


def test_exportacao_xlsx_com_linhas_completas():
    # Testar exportação a partir de um objeto Razao lido e conciliado
    conteudo = _criar_xlsx_exemplo_bytes()
    razao = ler_xlsx(conteudo, "exemplo.xlsx")
    resumo = conciliar(razao)
    dados = razao.como_dicionario()
    dados["resumo"] = resumo

    excel_bytes = gerar_xlsx_colorido(dados)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.max_row > 5


def test_exportacao_e_releitura_roundtrip():
    # Cria uma planilha, lê com ler_xlsx, concilia, exporta com gerar_xlsx_colorido
    # e lê novamente os bytes gerados com ler_xlsx.
    conteudo_inicial = _criar_xlsx_exemplo_bytes()
    razao_1 = ler_xlsx(conteudo_inicial, "inicial.xlsx")
    resumo_1 = conciliar(razao_1)

    dados_1 = razao_1.como_dicionario()
    dados_1["resumo"] = resumo_1

    bytes_exportados = gerar_xlsx_colorido(dados_1)
    razao_2 = ler_xlsx(bytes_exportados, "releitura.xlsx")

    assert razao_2.conta == razao_1.conta
    assert len(razao_2.lancamentos) == len(razao_1.lancamentos)
    assert [l.valor for l in razao_2.lancamentos] == [l.valor for l in razao_1.lancamentos]
    assert [l.tipo for l in razao_2.lancamentos] == [l.tipo for l in razao_1.lancamentos]

    resumo_2 = conciliar(razao_2)
    assert resumo_2["quitados"]["qtd"] == resumo_1["quitados"]["qtd"]
    assert resumo_2["cobertura_valor"] == resumo_1["cobertura_valor"]


def test_exportacao_status_ambiguo():
    razao = Razao(
        arquivo="teste_ambiguo.xlsx",
        empresa="EMPRESA TESTE",
        conta="123",
        lancamentos=[
            Lancamento(id=1, linha=5, tipo="C", valor=5000, data="2026-05-01", historico="COMPRA A", contrapartida="10", status="ambiguo", candidatos=[2, 3]),
            Lancamento(id=2, linha=6, tipo="D", valor=5000, data="2026-05-02", historico="PAGTO 1", contrapartida="10", status="ambiguo", candidatos=[1]),
            Lancamento(id=3, linha=7, tipo="D", valor=5000, data="2026-05-03", historico="PAGTO 2", contrapartida="10", status="ambiguo", candidatos=[1]),
        ],
        linhas=[
            {"tipo": "lancamento", "data": "2026-05-01", "historico": "COMPRA A", "contrapartida": "10", "debito": None, "credito": 1, "saldo": None},
            {"tipo": "lancamento", "data": "2026-05-02", "historico": "PAGTO 1", "contrapartida": "10", "debito": 2, "credito": None, "saldo": None},
            {"tipo": "lancamento", "data": "2026-05-03", "historico": "PAGTO 2", "contrapartida": "10", "debito": 3, "credito": None, "saldo": None},
        ],
    )
    dados = razao.como_dicionario()
    dados["resumo"] = {
        "debitos": {"qtd": 2, "valor": 10000},
        "creditos": {"qtd": 1, "valor": 5000},
        "quitados": {"qtd": 0, "valor": 0},
        "abertos": {"qtd": 0, "valor": 0},
        "sem_par": {"qtd": 0, "valor": 0},
        "ambiguos": {"qtd": 3, "valor": 15000},
        "cobertura_qtd": 0.0,
        "cobertura_valor": 0.0,
    }
    excel_bytes = gerar_xlsx_colorido(dados)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    # Encontra a linha com "Pendente (2 opções de par)"
    textos = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value]
    assert any("Pendente" in t for t in textos)

