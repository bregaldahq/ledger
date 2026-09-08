"""Exportação do Livro Razão conciliado para formato .xlsx com marcação visual.

Aplica estilos do Bregalda Design System:
- Quitado/Conciliado: Amarelo suave (#FEF9C3) com nota de amarração e dias.
- Em aberto: Destaque Lápis Vermelho (#FEE2E2).
- Sem par: Destaque Lápis Azul (#DBEAFE).
- Cabeçalhos institucionais Deep Ink (#2A2140) com texto em branco.
"""

from __future__ import annotations

import datetime
import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# Paleta de Cores
COR_DEEP_INK = "2A2140"
COR_BRANCO = "FFFFFF"
COR_CINZA_CLARO = "F3F4F6"
COR_BORDA = "E5E7EB"

COR_AMARELO_QUITADO = "FEF9C3"
COR_AMARELO_TEXTO = "713F12"

COR_VERMELHO_ABERTO = "FEE2E2"
COR_VERMELHO_TEXTO = "991B1B"

COR_AZUL_SEMPAR = "DBEAFE"
COR_AZUL_TEXTO = "1E40AF"

COR_AMBAR_AMBIGUO = "FEF3C7"
COR_AMBAR_TEXTO = "92400E"

FORMATO_MOEDA = '"R$" #,##0.00'
FORMATO_DATA = "DD/MM/YYYY"


def _formatar_data_para_celula(data_iso: str | None) -> tuple[datetime.date | str | None, str]:
    if not data_iso:
        return "", ""
    try:
        dt = datetime.date.fromisoformat(str(data_iso))
        return dt, FORMATO_DATA
    except Exception:
        # Se vier em formato texto ex: 01/07/2026
        return str(data_iso), ""


def _obter_estilo_status(lancamento: dict) -> tuple[str, PatternFill | None, Font]:
    status = str(lancamento.get("status", "aberto")).lower()
    par = lancamento.get("par")
    dias = lancamento.get("dias")

    if status in ("quitado", "conciliado"):
        fill = PatternFill(start_color=COR_AMARELO_QUITADO, end_color=COR_AMARELO_QUITADO, fill_type="solid")
        fonte = Font(name="Segoe UI", size=10, color=COR_AMARELO_TEXTO, bold=True)
        if dias is not None:
            if dias == 0:
                texto = f"Quitado no mesmo dia (par #{par})"
            elif dias == 1:
                texto = f"Quitado em 1 dia (par #{par})"
            else:
                texto = f"Quitado em {dias} dias (par #{par})"
        elif par is not None:
            texto = f"Quitado (par #{par})"
        else:
            texto = "Quitado"
        return texto, fill, fonte

    if status == "ambiguo":
        fill = PatternFill(start_color=COR_AMBAR_AMBIGUO, end_color=COR_AMBAR_AMBIGUO, fill_type="solid")
        fonte = Font(name="Segoe UI", size=10, color=COR_AMBAR_TEXTO, bold=True)
        candidatos = lancamento.get("candidatos") or []
        qtd_cand = len(candidatos)
        texto = f"Pendente ({qtd_cand} opções de par)" if qtd_cand > 0 else "Pendente de seleção"
        return texto, fill, fonte

    if status == "aberto":
        fill = PatternFill(start_color=COR_VERMELHO_ABERTO, end_color=COR_VERMELHO_ABERTO, fill_type="solid")
        fonte = Font(name="Segoe UI", size=10, color=COR_VERMELHO_TEXTO, bold=True)
        return "Em aberto", fill, fonte

    if status == "sem_par":
        fill = PatternFill(start_color=COR_AZUL_SEMPAR, end_color=COR_AZUL_SEMPAR, fill_type="solid")
        fonte = Font(name="Segoe UI", size=10, color=COR_AZUL_TEXTO, bold=True)
        return "Sem par no mês", fill, fonte

    return status, None, Font(name="Segoe UI", size=10)


def gerar_xlsx_colorido(razao_dict: dict) -> bytes:
    """Gera uma pasta de trabalho Excel (.xlsx) colorida a partir do razão conciliado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Razão Conciliado"
    ws.views.sheetView[0].showGridLines = True

    borda_fina = Border(
        left=Side(style="thin", color=COR_BORDA),
        right=Side(style="thin", color=COR_BORDA),
        top=Side(style="thin", color=COR_BORDA),
        bottom=Side(style="thin", color=COR_BORDA),
    )

    fonte_titulo = Font(name="Segoe UI", size=13, bold=True, color=COR_DEEP_INK)
    fonte_meta_rotulo = Font(name="Segoe UI", size=10, bold=True, color=COR_DEEP_INK)
    fonte_meta_valor = Font(name="Segoe UI", size=10, color="374151")
    fonte_cabecalho = Font(name="Segoe UI", size=10, bold=True, color=COR_BRANCO)
    fill_cabecalho = PatternFill(start_color=COR_DEEP_INK, end_color=COR_DEEP_INK, fill_type="solid")

    # ---- 1. Cabeçalho Institucional
    ws.cell(1, 1, "CONFERÊNCIA DE LIVRO RAZÃO CONTÁBIL").font = fonte_titulo

    metadados = [
        ("Empresa:", razao_dict.get("empresa", "")),
        ("CNPJ:", razao_dict.get("cnpj", "")),
        ("Período:", razao_dict.get("periodo", "")),
        (
            "Conta:",
            f"{razao_dict.get('conta', '')} - {razao_dict.get('conta_nome', '')}".strip(" -"),
        ),
    ]

    linha_atual = 3
    for rotulo, valor in metadados:
        cel_r = ws.cell(linha_atual, 1, rotulo)
        cel_r.font = fonte_meta_rotulo
        cel_v = ws.cell(linha_atual, 2, valor)
        cel_v.font = fonte_meta_valor
        linha_atual += 1

    linha_atual += 1  # Espaço antes da tabela

    # ---- 2. Cabeçalho da Tabela
    colunas_titulos = [
        "Data",
        "Histórico",
        "Cta.C.Part.",
        "Débito",
        "Crédito",
        "Saldo",
        "Status / Conciliação",
    ]

    linha_titulos = linha_atual
    for col_idx, titulo in enumerate(colunas_titulos, start=1):
        c = ws.cell(linha_titulos, col_idx, titulo)
        c.font = fonte_cabecalho
        c.fill = fill_cabecalho
        c.border = borda_fina
        c.alignment = Alignment(
            horizontal="center" if col_idx in (1, 7) else ("right" if col_idx in (4, 5, 6) else "left"),
            vertical="center",
        )

    linha_atual += 1
    linha_dados_inicio = linha_atual

    # ---- 3. Dados dos Lançamentos
    lancamentos = razao_dict.get("lancamentos", [])
    lancamentos_map = {l["id"]: l for l in lancamentos}
    linhas = razao_dict.get("linhas", [])

    # Se há 'linhas' estruturadas preservadas do leitor, priorizamos o layout original
    linhas_para_processar = []
    if linhas:
        for item in linhas:
            tipo = item.get("tipo")
            if tipo == "conta":
                linhas_para_processar.append({"tipo": "conta", "item": item})
            elif tipo == "saldo_anterior":
                linhas_para_processar.append({"tipo": "saldo_anterior", "item": item})
            elif tipo == "total":
                linhas_para_processar.append({"tipo": "total", "item": item})
            elif tipo == "lancamento":
                deb_id = item.get("debito")
                cred_id = item.get("credito")
                if deb_id is not None and cred_id is not None:
                    # Ambas as colunas na mesma linha
                    linhas_para_processar.append({"tipo": "lancamento_duplo", "deb": lancamentos_map.get(deb_id), "cred": lancamentos_map.get(cred_id), "item": item})
                elif deb_id is not None:
                    linhas_para_processar.append({"tipo": "lancamento", "lanc": lancamentos_map.get(deb_id), "item": item})
                elif cred_id is not None:
                    linhas_para_processar.append({"tipo": "lancamento", "lanc": lancamentos_map.get(cred_id), "item": item})
    else:
        # Se não há 'linhas' salvas, exportamos a partir de 'saldo_anterior', 'lancamentos' e 'totais'
        if razao_dict.get("saldo_anterior") is not None:
            linhas_para_processar.append({"tipo": "saldo_anterior", "item": {"saldo": razao_dict["saldo_anterior"]}})
        for l in lancamentos:
            linhas_para_processar.append({"tipo": "lancamento", "lanc": l, "item": {}})
        if razao_dict.get("total_debito") is not None or razao_dict.get("total_credito") is not None:
            linhas_para_processar.append({
                "tipo": "total",
                "item": {
                    "debito": razao_dict.get("total_debito"),
                    "credito": razao_dict.get("total_credito"),
                },
            })

    for entrada in linhas_para_processar:
        tipo = entrada["tipo"]

        if tipo == "conta":
            item = entrada["item"]
            codigo = item.get("codigo", "")
            nome = item.get("nome", "")
            ws.cell(linha_atual, 1, "")
            c = ws.cell(linha_atual, 2, f"Conta: {codigo}  {nome}".strip())
            c.font = Font(name="Segoe UI", size=10, bold=True, color=COR_DEEP_INK)
            fill_conta = PatternFill(start_color=COR_CINZA_CLARO, end_color=COR_CINZA_CLARO, fill_type="solid")
            for c_idx in range(1, 8):
                ws.cell(linha_atual, c_idx).fill = fill_conta
                ws.cell(linha_atual, c_idx).border = borda_fina
            linha_atual += 1

        elif tipo == "saldo_anterior":
            item = entrada["item"]
            saldo_cents = item.get("saldo")
            ws.cell(linha_atual, 1, "")
            c_hist = ws.cell(linha_atual, 2, "SALDO ANTERIOR")
            c_hist.font = Font(name="Segoe UI", size=10, bold=True)
            if saldo_cents is not None:
                c_saldo = ws.cell(linha_atual, 6, saldo_cents / 100.0)
                c_saldo.number_format = FORMATO_MOEDA
                c_saldo.font = Font(name="Segoe UI", size=10, bold=True)
                c_saldo.alignment = Alignment(horizontal="right", vertical="center")
            for c_idx in range(1, 8):
                ws.cell(linha_atual, c_idx).border = borda_fina
            linha_atual += 1

        elif tipo == "total":
            item = entrada["item"]
            deb_cents = item.get("debito")
            cred_cents = item.get("credito")
            ws.cell(linha_atual, 1, "")
            c_hist = ws.cell(linha_atual, 2, "Total do mês")
            c_hist.font = Font(name="Segoe UI", size=10, bold=True)
            if deb_cents is not None:
                c_deb = ws.cell(linha_atual, 4, deb_cents / 100.0)
                c_deb.number_format = FORMATO_MOEDA
                c_deb.font = Font(name="Segoe UI", size=10, bold=True)
                c_deb.alignment = Alignment(horizontal="right", vertical="center")
            if cred_cents is not None:
                c_cred = ws.cell(linha_atual, 5, cred_cents / 100.0)
                c_cred.number_format = FORMATO_MOEDA
                c_cred.font = Font(name="Segoe UI", size=10, bold=True)
                c_cred.alignment = Alignment(horizontal="right", vertical="center")
            fill_total = PatternFill(start_color=COR_CINZA_CLARO, end_color=COR_CINZA_CLARO, fill_type="solid")
            for c_idx in range(1, 8):
                ws.cell(linha_atual, c_idx).fill = fill_total
                ws.cell(linha_atual, c_idx).border = borda_fina
            linha_atual += 1

        elif tipo == "lancamento":
            l = entrada["lanc"]
            if not l:
                continue
            item = entrada.get("item", {})

            data_val, num_fmt = _formatar_data_para_celula(l.get("data"))
            historico = l.get("historico", "")
            contrapartida = l.get("contrapartida", "")
            tipo_mov = l.get("tipo")
            valor = l.get("valor", 0) / 100.0
            saldo_cents = item.get("saldo", l.get("saldo"))

            texto_status, fill_status, fonte_status = _obter_estilo_status(l)

            # Data
            c_data = ws.cell(linha_atual, 1, data_val)
            if num_fmt:
                c_data.number_format = num_fmt
            c_data.alignment = Alignment(horizontal="center", vertical="center")

            # Histórico
            c_hist = ws.cell(linha_atual, 2, historico)
            c_hist.alignment = Alignment(horizontal="left", vertical="center")

            # Contrapartida
            c_ctap = ws.cell(linha_atual, 3, contrapartida)
            c_ctap.alignment = Alignment(horizontal="left", vertical="center")

            # Débito
            c_deb = ws.cell(linha_atual, 4)
            if tipo_mov == "D":
                c_deb.value = valor
                c_deb.number_format = FORMATO_MOEDA
            c_deb.alignment = Alignment(horizontal="right", vertical="center")

            # Crédito
            c_cred = ws.cell(linha_atual, 5)
            if tipo_mov == "C":
                c_cred.value = valor
                c_cred.number_format = FORMATO_MOEDA
            c_cred.alignment = Alignment(horizontal="right", vertical="center")

            # Saldo
            c_saldo = ws.cell(linha_atual, 6)
            if saldo_cents is not None:
                c_saldo.value = saldo_cents / 100.0
                c_saldo.number_format = FORMATO_MOEDA
            c_saldo.alignment = Alignment(horizontal="right", vertical="center")

            # Status / Conciliação
            c_stat = ws.cell(linha_atual, 7, texto_status)
            c_stat.font = fonte_status
            c_stat.alignment = Alignment(horizontal="center", vertical="center")

            # Aplicação do preenchimento colorido e bordas na linha
            for col_idx in range(1, 8):
                cell = ws.cell(linha_atual, col_idx)
                cell.border = borda_fina
                if fill_status:
                    cell.fill = fill_status
                if col_idx != 7 and not cell.font.bold:
                    cell.font = Font(name="Segoe UI", size=10)

            linha_atual += 1

        elif tipo == "lancamento_duplo":
            # Caso especial com débito e crédito na mesma linha
            l_deb = entrada["deb"]
            l_cred = entrada["cred"]
            item = entrada["item"]

            data_val, num_fmt = _formatar_data_para_celula(item.get("data") or (l_cred.get("data") if l_cred else None))
            historico = item.get("historico", "")
            contrapartida = item.get("contrapartida", "")
            saldo_cents = item.get("saldo")

            c_data = ws.cell(linha_atual, 1, data_val)
            if num_fmt:
                c_data.number_format = num_fmt
            c_data.alignment = Alignment(horizontal="center", vertical="center")

            ws.cell(linha_atual, 2, historico).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(linha_atual, 3, contrapartida).alignment = Alignment(horizontal="left", vertical="center")

            c_deb = ws.cell(linha_atual, 4)
            if l_deb:
                c_deb.value = l_deb.get("valor", 0) / 100.0
                c_deb.number_format = FORMATO_MOEDA
            c_deb.alignment = Alignment(horizontal="right", vertical="center")

            c_cred = ws.cell(linha_atual, 5)
            if l_cred:
                c_cred.value = l_cred.get("valor", 0) / 100.0
                c_cred.number_format = FORMATO_MOEDA
            c_cred.alignment = Alignment(horizontal="right", vertical="center")

            c_saldo = ws.cell(linha_atual, 6)
            if saldo_cents is not None:
                c_saldo.value = saldo_cents / 100.0
                c_saldo.number_format = FORMATO_MOEDA
            c_saldo.alignment = Alignment(horizontal="right", vertical="center")

            # Status combinado
            partes_status = []
            preenchimento = None
            if l_cred:
                txt, fill_c, _ = _obter_estilo_status(l_cred)
                partes_status.append(f"C: {txt}")
                preenchimento = fill_c
            if l_deb:
                txt, fill_d, _ = _obter_estilo_status(l_deb)
                partes_status.append(f"D: {txt}")
                if not preenchimento:
                    preenchimento = fill_d

            c_stat = ws.cell(linha_atual, 7, " | ".join(partes_status))
            c_stat.alignment = Alignment(horizontal="center", vertical="center")
            c_stat.font = Font(name="Segoe UI", size=9, bold=True)

            for col_idx in range(1, 8):
                cell = ws.cell(linha_atual, col_idx)
                cell.border = borda_fina
                if preenchimento:
                    cell.fill = preenchimento

            linha_atual += 1

    # ---- 4. Resumo da Conciliação no rodapé (se disponível)
    resumo = razao_dict.get("resumo")
    if resumo:
        linha_atual += 1
        ws.cell(linha_atual, 1, "RESUMO DA CONCILIAÇÃO").font = Font(name="Segoe UI", size=11, bold=True, color=COR_DEEP_INK)
        linha_atual += 1

        cobertura = resumo.get("cobertura_valor", 0.0)
        quitados = resumo.get("quitados", {})
        abertos = resumo.get("abertos", {})
        sem_par = resumo.get("sem_par", {})

        itens_resumo = [
            ("Cobertura do Período:", f"{cobertura:.1f}%"),
            ("Aquisições Quitadas:", f"{quitados.get('qtd', 0)} ({_formatar_moeda_br(quitados.get('valor', 0))})"),
            ("Aquisições em Aberto:", f"{abertos.get('qtd', 0)} ({_formatar_moeda_br(abertos.get('valor', 0))})"),
            ("Débitos Sem Par:", f"{sem_par.get('qtd', 0)} ({_formatar_moeda_br(sem_par.get('valor', 0))})"),
        ]

        for rotulo, val in itens_resumo:
            c_r = ws.cell(linha_atual, 1, rotulo)
            c_r.font = Font(name="Segoe UI", size=9, bold=True, color="4B5563")
            c_v = ws.cell(linha_atual, 2, val)
            c_v.font = Font(name="Segoe UI", size=9, bold=True, color=COR_DEEP_INK)
            linha_atual += 1

    # ---- 5. Auto-ajuste de largura das colunas
    larguras_base = {
        1: 14,  # Data
        2: 44,  # Histórico
        3: 16,  # Cta.C.Part.
        4: 18,  # Débito
        5: 18,  # Crédito
        6: 18,  # Saldo
        7: 35,  # Status / Conciliação
    }

    for col_idx in range(1, 8):
        col_letra = get_column_letter(col_idx)
        max_len = larguras_base.get(col_idx, 15)
        for r in range(linha_dados_inicio - 1, ws.max_row + 1):
            val = ws.cell(r, col_idx).value
            if val:
                max_len = max(max_len, min(len(str(val)) + 3, 65))
        ws.column_dimensions[col_letra].width = max_len

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _formatar_moeda_br(centavos: int) -> str:
    valor = centavos / 100.0
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
